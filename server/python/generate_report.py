#!/usr/bin/env python3
"""
generate_report.py — Assemble the Macquarie submission package from a JSON job.
Payload-driven port of the original cowork pipeline (no workspace/folders).

Usage:  python3 generate_report.py job.json
job.json: {
  profile, invoiceNumber, periodLabel, generatedDate,
  expenses: [{date, description, amountVND, accountCode, notes, ...}],
  receiptFiles: ["/abs/path", ...]  (B&W scans, in expense order),
  outDir, invoiceTemplate
}
Prints (last line): {"combinedPdf": "<path>", "expenseXlsx": "<path>"}

Output order in the combined PDF: invoice -> expense detail -> merged receipts.
Nothing is written outside outDir; the server deletes outDir afterwards.
"""
import sys, os, json, shutil, subprocess, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Expense detail XLSX (layout preserved from the cowork pipeline) ──────────
def generate_expense_xlsx(rows, out_path, period_label, base_currency="VND"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    HEADER_BG, ALT, TOTAL_BG = "1F4E79", "EBF3FB", "BDD7EE"
    thin = Side(style="thin", color="BFBFBF")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr(col, val, w=None):
        c = ws.cell(row=1, column=col, value=val)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bdr
        if w:
            ws.column_dimensions[get_column_letter(col)].width = w

    hdr(1, "#", 8); hdr(2, "Date", 11); hdr(3, "Description", 44)
    hdr(4, "Account Code", 18); hdr(5, f"Amount ({base_currency})", 16)
    ws.row_dimensions[1].height = 24

    def parse_date(d):
        try:
            return datetime.datetime.strptime(d, "%Y-%m-%d").date()
        except Exception:
            return d

    sorted_rows = sorted(rows, key=lambda x: x.get("date", ""))
    excel_row, receipt_no, subtotals = 2, 1, {}
    for r in sorted_rows:
        acc = r.get("accountCode", "")
        amt = int(round(float(r.get("amountVND", 0) or 0)))
        subtotals[acc] = subtotals.get(acc, 0) + amt
        fill = ALT if excel_row % 2 == 0 else "FFFFFF"
        vals = [receipt_no, parse_date(r.get("date", "")), r.get("description", ""), acc, amt]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=excel_row, column=c, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = bdr
            cell.alignment = Alignment(vertical="center")
            if c == 2:
                cell.number_format = "DD/MM/YYYY"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c == 5:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
        excel_row += 1
        receipt_no += 1

    for acc in sorted(subtotals):
        ws.merge_cells(f"A{excel_row}:D{excel_row}")
        lbl = ws.cell(row=excel_row, column=1, value=f"Subtotal - {acc}")
        lbl.font = Font(name="Arial", bold=True, size=10)
        lbl.fill = PatternFill("solid", fgColor=TOTAL_BG)
        lbl.alignment = Alignment(horizontal="right", vertical="center")
        lbl.border = bdr
        val = ws.cell(row=excel_row, column=5, value=subtotals[acc])
        val.font = Font(name="Arial", bold=True, size=10)
        val.fill = PatternFill("solid", fgColor=TOTAL_BG)
        val.number_format = "#,##0"
        val.alignment = Alignment(horizontal="right", vertical="center")
        val.border = bdr
        excel_row += 1

    ws.merge_cells(f"A{excel_row}:D{excel_row}")
    gl = ws.cell(row=excel_row, column=1, value="GRAND TOTAL")
    gl.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    gl.fill = PatternFill("solid", fgColor="1F4E79")
    gl.alignment = Alignment(horizontal="right", vertical="center")
    gl.border = bdr
    grand = sum(subtotals.values())
    gv = ws.cell(row=excel_row, column=5, value=grand)
    gv.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    gv.fill = PatternFill("solid", fgColor="1F4E79")
    gv.number_format = "#,##0"
    gv.alignment = Alignment(horizontal="right", vertical="center")
    gv.border = bdr

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = 9
    wb.save(out_path)
    return grand


# ── Invoice XLSX — fill the bundled template from the profile ────────────────
def generate_invoice(job, total, out_path):
    tmpl = job["invoiceTemplate"]
    if not os.path.exists(tmpl):
        raise FileNotFoundError(f"invoice template missing: {tmpl}")
    wb = openpyxl.load_workbook(tmpl)
    ws = wb[wb.sheetnames[0]]
    p = job.get("profile", {})
    sub = p.get("submitter", {}) or {}
    bank = p.get("bank", {}) or {}
    to = p.get("invoiceTo", {}) or {}

    def setc(coord, value, fmt=None):
        if value is None or value == "":
            return
        ws[coord] = value
        if fmt:
            ws[coord].number_format = fmt

    # From / submitter block
    setc("A4", sub.get("name"))
    setc("A5", sub.get("addressLine1"))
    setc("A6", sub.get("addressLine2"))
    setc("A7", sub.get("country"))
    setc("A8", sub.get("phone"))
    # Invoice meta
    try:
        gdate = datetime.datetime.strptime(job.get("generatedDate", ""), "%Y-%m-%d").date()
    except Exception:
        gdate = datetime.date.today()
    setc("D4", gdate, "DD/MM/YYYY")
    setc("D6", job["invoiceNumber"])
    setc("D8", p.get("vendorId"))
    # Invoice-to block
    setc("A11", to.get("name"))
    addr_lines = [l for l in (to.get("address", "") or "").split("\n") if l.strip()]
    for i, line in enumerate(addr_lines[:3]):
        setc(f"A{12 + i}", line)
    setc("A15", to.get("email"))
    # Line item + total
    setc("A18", f"Expenses Reimbursement for {job.get('periodLabel', '')}")
    setc("D18", total, "#,##0")
    # Bank block
    setc("B33", bank.get("accountName"))
    setc("B34", bank.get("accountNumber"))
    setc("B35", bank.get("swift"))
    setc("B36", bank.get("bankName"))

    wb.save(out_path)
    return out_path


# ── xlsx -> pdf via LibreOffice headless ─────────────────────────────────────
def find_soffice():
    for c in ["soffice", "libreoffice", "/opt/homebrew/bin/soffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/bin/soffice", "/usr/bin/libreoffice"]:
        if shutil.which(c) or os.path.exists(c):
            return c
    return None


def xlsx_to_pdf(xlsx_path, out_dir):
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not found — required for PDF export")
    res = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(xlsx_path)],
                         capture_output=True, text=True, timeout=90)
    pdf = Path(out_dir) / (Path(xlsx_path).stem + ".pdf")
    if res.returncode != 0 or not pdf.exists():
        raise RuntimeError(f"LibreOffice conversion failed: {res.stderr.strip()[:300]}")
    return pdf


# ── Receipt merge (pdf passthrough + image -> pdf) ───────────────────────────
def merge_receipts(receipt_files, out_path):
    from pypdf import PdfWriter, PdfReader
    from PIL import Image
    import io
    pages = []
    for fp in receipt_files:
        fp = Path(fp)
        if not fp.exists() or fp.stat().st_size == 0:
            continue
        try:
            if fp.suffix.lower() == ".pdf" or _looks_pdf(fp):
                pages.append(fp.read_bytes())
            else:
                img = Image.open(fp).convert("RGB")
                buf = io.BytesIO()
                img.save(buf, format="PDF")
                pages.append(buf.getvalue())
        except Exception:
            continue
    if not pages:
        return None
    writer = PdfWriter()
    for b in pages:
        for pg in PdfReader(io.BytesIO(b)).pages:
            writer.add_page(pg)
    with open(out_path, "wb") as f:
        writer.write(f)
    return out_path


def _looks_pdf(fp):
    try:
        with open(fp, "rb") as f:
            return f.read(5) == b"%PDF-"
    except Exception:
        return False


def combine(invoice_pdf, expense_pdf, receipts_pdf, out_path):
    from pypdf import PdfWriter, PdfReader
    writer = PdfWriter()
    for src in [invoice_pdf, expense_pdf, receipts_pdf]:
        if src and Path(src).exists():
            for pg in PdfReader(str(src)).pages:
                writer.add_page(pg)
    with open(out_path, "wb") as f:
        writer.write(f)
    return out_path


def main():
    job = json.loads(Path(sys.argv[1]).read_text())
    out = Path(job["outDir"])
    expenses = job["expenses"]

    expense_xlsx = out / "expenses.xlsx"
    total = generate_expense_xlsx(expenses, expense_xlsx, job.get("periodLabel", ""), job.get("baseCurrency", "VND"))

    invoice_xlsx = out / "invoice.xlsx"
    generate_invoice(job, total, invoice_xlsx)

    invoice_pdf = xlsx_to_pdf(invoice_xlsx, out)
    expense_pdf = xlsx_to_pdf(expense_xlsx, out)
    receipts_pdf = merge_receipts(job.get("receiptFiles", []), out / "receipts.pdf")

    combined = combine(invoice_pdf, expense_pdf, receipts_pdf, out / "combined.pdf")
    print(json.dumps({"combinedPdf": str(combined), "expenseXlsx": str(expense_xlsx)}))


if __name__ == "__main__":
    main()
